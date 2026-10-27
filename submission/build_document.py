from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
out=Path(__file__).parent
D=Document(); S=D.sections[0]
S.top_margin=S.bottom_margin=Inches(.7); S.left_margin=S.right_margin=Inches(.8)
S.page_width=Inches(8.5);S.page_height=Inches(11)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=D.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)

for st in D.styles:
 for border in list(st.element.xpath('.//w:pBdr')):
  border.getparent().remove(border)
normal=D.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
D.styles['Title'].font.size=Pt(27); D.styles['Title'].paragraph_format.space_after=Pt(8)
D.styles['Heading 1'].font.size=Pt(13);D.styles['Heading 1'].paragraph_format.space_before=Pt(11);D.styles['Heading 1'].paragraph_format.space_after=Pt(5)
D.add_heading('Quiet Classroom',0)
D.add_paragraph('Interactive Web Design Reflection',style='Subtitle')
D.add_heading('Idea and design intentions',1)
D.add_paragraph('Quiet Classroom is an interactive website that encourages a quieter study environment through a sleeping fish. Instead of displaying only a warning or a number, the page turns classroom sound into a small story: staying quiet lets the fish rest, while sustained noise wakes it up. The intention is to make self regulation feel gentle, playful, and easy to understand.')
D.add_paragraph('The interface pairs an illustrated room and aquarium with a study control panel. Soft gradients, rounded shapes, translucent surfaces, and warm furniture create a calm atmosphere. The window and background change with the local time of day. A live sound meter, quiet time, score, and session summary give feedback. Manual, duration, and scheduled modes support different study routines, while a sensitivity slider lets users adjust the response to their environment.')
D.add_heading('How I used AI tools',1)
D.add_paragraph('I used ChatGPT and Codex to help develop and revise the website. My requests focused on an interactive Quiet Classroom experience and, in this revision, a more lively fish, clearer visual design, and a review of unreasonable behavior or code bugs. I also asked AI to check whether the repository contained CSS and JavaScript and to help prepare this reflection.')
D.add_paragraph('AI provided the HTML interface, CSS room and fish artwork, and JavaScript for microphone input, sound level feedback, timing, scoring, and the result dialog. The original repository stored its CSS and JavaScript inside the HTML rather than as separate files. During this revision, AI extracted them into style.css and script.js and updated the HTML links so the structure, presentation, and behavior are easier to inspect and maintain.')
D.add_paragraph('My contribution was to set the concept, choose the quiet study scenario, and direct the next changes: improve the fish character and review the interaction logic. I used AI as an implementation assistant and requested another review rather than treating the generated code as finished. The revision added breathing, fin and tail animation to the fish and improved the panel hierarchy. AI also corrected microphone cleanup, cancellation during startup, repeated noise counting, and simulated sound scoring. I requested these refinements so the visual idea and the session behavior work together.')
D.add_heading('HTML and CSS work',1)
p=D.add_paragraph('GitHub repository: ')
url='https://github.com/yz5099-sys/Quiet-Classroom'
rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid);r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'176B69');pr.append(c);r.append(pr);t=OxmlElement('w:t');t.text=url;r.append(t);h.append(r);p._p.append(h)
D.add_paragraph('Files: index.html for structure, style.css for visual design, and script.js for interaction. The sound percentage is a relative microphone level, not a calibrated decibel measurement.')
D.core_properties.title='Quiet Classroom Interactive Web Design Reflection';D.core_properties.author='';D.core_properties.subject='Design intentions and AI assisted development'
D.save(out/'Quiet_Classroom_Design_Reflection.docx')
